"""Check the generated synopsis against the PSIT formatting rules."""

from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

PATH = Path(__file__).resolve().parents[1] / "CreatorLens_Project_Synopsis.docx"
doc = Document(str(PATH))
problems: list[str] = []

for index, section in enumerate(doc.sections):
    margins = [round(value.cm, 2) for value in (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin)]
    print(f"section {index}: margins (T,B,L,R) cm = {margins}; page {section.page_width.cm:.2f} x {section.page_height.cm:.2f} cm")
    if any(abs(margin - 2.54) > 0.01 for margin in margins):
        problems.append(f"section {index} margins {margins}")
footer = doc.sections[-1].footer
footer_xml = footer._element.xml
print("body footer: PAGE field =", "PAGE" in footer_xml, "| alignment =", footer.paragraphs[0].alignment)
if "PAGE" not in footer_xml or footer.paragraphs[0].alignment != WD_ALIGN_PARAGRAPH.RIGHT:
    problems.append("page number field missing or not right-aligned")
print("cover section has footer reference:", doc.sections[0]._sectPr.find(qn("w:footerReference")) is not None)

fonts = Counter()
sizes = Counter()
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if run.text.strip():
            fonts[run.font.name] += 1
            sizes[run.font.size.pt if run.font.size else None] += 1
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip():
                        fonts[run.font.name] += 1
print("run fonts:", dict(fonts))
print("run sizes (body paragraphs):", dict(sizes))
if set(fonts) != {"Times New Roman"}:
    problems.append(f"non-TNR runs: {dict(fonts)}")

texts = [paragraph.text for paragraph in doc.paragraphs]
start = texts.index("1. Introduction")
body = doc.paragraphs[start:]
justified = sum(1 for p in body if p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY)
spacing = Counter(p.paragraph_format.line_spacing for p in body if p.text.strip())
print("body paragraphs:", len(body), "| justified:", justified, "| line spacing values:", dict(spacing))

headings = [p.text for p in body if p.runs and p.runs[0].font.size and p.runs[0].font.size.pt == 14]
print("14 pt headings:", headings)
expected = ["1. Introduction", "2. Project Objective", "3. Feasibility Study", "4. Methodology / Planning of Work",
            "5. Tools/Technology Used", "6. References"]
if headings != expected:
    problems.append(f"headings {headings}")
subs = [p.text for p in body if p.runs and p.runs[0].bold and p.runs[0].font.size and p.runs[0].font.size.pt == 12
        and re.match(r"^\d\.\d ", p.text)]
print("numbered subheadings:", subs)

images = len(doc.inline_shapes)
print("inline images:", images)
captions = [t for t in texts if t.startswith("Figure ") or t.startswith("Table ")]
print("captions:", captions)

full = "\n".join(texts) + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
ref_start = texts.index("6. References")
ref_entries = [t for t in texts[ref_start + 1:] if re.match(r"^\[\d+\]", t)]
cited = {int(n) for n in re.findall(r"\[(\d+)\]", "\n".join(texts[start:ref_start]))}
print("references listed:", len(ref_entries), "| cited numbers:", sorted(cited))
if cited != set(range(1, len(ref_entries) + 1)):
    problems.append("citation numbers do not match the reference list")
print("placeholders:", sorted(set(re.findall(r"\[[^\]]*PLACEHOLDER[^\]]*\]", full))))
print("cover lines:", [t for t in texts[:start] if t.strip()])
print("PROBLEMS:", problems or "none")
