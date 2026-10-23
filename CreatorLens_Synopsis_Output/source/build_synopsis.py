"""Build CreatorLens_Project_Synopsis.docx from the official PSIT synopsis template.

The template file is opened read-only; the output is saved as a new file.
Run from the project root:
    python CreatorLens_Synopsis_Output/source/build_synopsis.py
"""

from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

TEMPLATE = Path(r"C:\Users\anshi\Downloads\CSE_Mini Project Synopsis_Format 2026_27.docx")
OUT_DIR = Path(__file__).resolve().parents[1]
OUTPUT = OUT_DIR / "CreatorLens_Project_Synopsis.docx"
FONT = "Times New Roman"

REFERENCES: dict[str, str] = {
    "react": 'Meta Open Source, "React documentation." [Online]. Available: https://react.dev/ (accessed Oct. 21, 2026).',
    "typescript": 'Microsoft, "TypeScript documentation." [Online]. Available: https://www.typescriptlang.org/docs/ (accessed Oct. 21, 2026).',
    "vite": 'Vite, "Vite guide." [Online]. Available: https://vite.dev/guide/ (accessed Oct. 21, 2026).',
    "fastapi": 'S. Ramírez, "FastAPI documentation." [Online]. Available: https://fastapi.tiangolo.com/ (accessed Oct. 21, 2026).',
    "pandas": 'W. McKinney, "Data structures for statistical computing in Python," in Proc. 9th Python in Science Conf. (SciPy 2010), Austin, TX, USA, 2010, pp. 56–61, doi: 10.25080/Majora-92bf1922-00a.',
    "numpy": 'C. R. Harris et al., "Array programming with NumPy," Nature, vol. 585, no. 7825, pp. 357–362, Sep. 2020, doi: 10.1038/s41586-020-2649-2.',
    "sqlite": 'SQLite, "SQLite documentation." [Online]. Available: https://www.sqlite.org/docs.html (accessed Oct. 21, 2026).',
    "ytdlp": 'yt-dlp contributors, "yt-dlp: A feature-rich command-line audio/video downloader," GitHub repository. [Online]. Available: https://github.com/yt-dlp/yt-dlp (accessed Oct. 21, 2026).',
    "sklearn": 'F. Pedregosa et al., "Scikit-learn: Machine learning in Python," J. Mach. Learn. Res., vol. 12, pp. 2825–2830, 2011.',
    "xgboost": 'T. Chen and C. Guestrin, "XGBoost: A scalable tree boosting system," in Proc. 22nd ACM SIGKDD Int. Conf. Knowledge Discovery and Data Mining (KDD \'16), San Francisco, CA, USA, 2016, pp. 785–794, doi: 10.1145/2939672.2939785.',
    "holt": 'C. C. Holt, "Forecasting seasonals and trends by exponentially weighted moving averages," Int. J. Forecasting, vol. 20, no. 1, pp. 5–10, 2004 (reprint of 1957 ONR memorandum), doi: 10.1016/j.ijforecast.2003.09.015.',
    "statsmodels": 'S. Seabold and J. Perktold, "Statsmodels: Econometric and statistical modeling with Python," in Proc. 9th Python in Science Conf. (SciPy 2010), Austin, TX, USA, 2010, pp. 92–96, doi: 10.25080/Majora-92bf1922-011.',
    "gemini": 'Google, "Gemini API documentation." [Online]. Available: https://ai.google.dev/gemini-api/docs (accessed Oct. 21, 2026).',
    "pytest": 'pytest-dev, "pytest documentation." [Online]. Available: https://docs.pytest.org/ (accessed Oct. 21, 2026).',
    "youtube_tos": 'YouTube, "Terms of Service." [Online]. Available: https://www.youtube.com/t/terms (accessed Oct. 21, 2026).',
    "rfc8018": 'K. Moriarty, B. Kaliski, and A. Rusch, "PKCS #5: Password-based cryptography specification version 2.1," Internet Engineering Task Force, RFC 8018, Jan. 2017, doi: 10.17487/RFC8018.',
    "owasp": 'OWASP Foundation, "Password storage cheat sheet," OWASP Cheat Sheet Series. [Online]. Available: https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html (accessed Oct. 21, 2026).',
    "dpdp": 'Government of India, "The Digital Personal Data Protection Act, 2023," Act No. 22 of 2023, The Gazette of India, Aug. 2023.',
    "cart": 'L. Breiman, J. H. Friedman, R. A. Olshen, and C. J. Stone, Classification and Regression Trees. Belmont, CA, USA: Wadsworth, 1984.',
    "rf": 'L. Breiman, "Random forests," Machine Learning, vol. 45, no. 1, pp. 5–32, Oct. 2001, doi: 10.1023/A:1010933404324.',
    "treeshap": 'S. M. Lundberg et al., "From local explanations to global understanding with explainable AI for trees," Nature Machine Intelligence, vol. 2, no. 1, pp. 56–67, Jan. 2020, doi: 10.1038/s42256-019-0138-9.',
    "fpp": 'R. J. Hyndman and G. Athanasopoulos, Forecasting: Principles and Practice, 3rd ed. Melbourne, Australia: OTexts, 2021. [Online]. Available: https://otexts.com/fpp3/ (accessed Oct. 21, 2026).',
    "tashman": 'L. J. Tashman, "Out-of-sample tests of forecasting accuracy: An analysis and review," Int. J. Forecasting, vol. 16, no. 4, pp. 437–450, Oct. 2000, doi: 10.1016/S0169-2070(00)00065-0.',
    "recharts": 'Recharts, "Recharts: A composable charting library built on React components." [Online]. Available: https://recharts.org/ (accessed Oct. 21, 2026).',
    "tailwind": 'Tailwind Labs, "Tailwind CSS documentation." [Online]. Available: https://tailwindcss.com/docs (accessed Oct. 21, 2026).',
}
ORDER: list[str] = []


def cite(*keys: str) -> str:
    numbers = []
    for key in keys:
        if key not in REFERENCES:
            raise KeyError(key)
        if key not in ORDER:
            ORDER.append(key)
        numbers.append(ORDER.index(key) + 1)
    return "[" + "], [".join(str(number) for number in numbers) + "]"


def set_run_font(run, size=None, bold=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), FONT)
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(attribute)) is not None:
            del fonts.attrib[qn(attribute)]
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def fmt(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, spacing=1.5, keep=False, indent=None, first=None):
    pf = paragraph.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing
    pf.keep_with_next = keep
    if indent is not None:
        pf.left_indent = Cm(indent)
    if first is not None:
        pf.first_line_indent = Cm(first)
    return paragraph


class Builder:
    def __init__(self, doc: Document):
        self.doc = doc
        self.figure = {}
        self.table_count = {}

    def outline(self, paragraph, level):
        ppr = paragraph._p.get_or_add_pPr()
        node = OxmlElement("w:outlineLvl")
        node.set(qn("w:val"), str(level))
        ppr.append(node)

    def heading(self, text, page_break=False):
        paragraph = self.doc.add_paragraph()
        if page_break:
            paragraph.paragraph_format.page_break_before = True
        run = paragraph.add_run(text)
        set_run_font(run, 14, bold=True)
        fmt(paragraph, WD_ALIGN_PARAGRAPH.LEFT, before=12, after=6, keep=True)
        self.outline(paragraph, 0)
        return paragraph

    def subheading(self, text):
        paragraph = self.doc.add_paragraph()
        run = paragraph.add_run(text)
        set_run_font(run, 12, bold=True)
        fmt(paragraph, WD_ALIGN_PARAGRAPH.LEFT, before=8, after=4, keep=True)
        self.outline(paragraph, 1)
        return paragraph

    def para(self, *segments, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep=False, indent=None, first=None):
        paragraph = self.doc.add_paragraph()
        for segment in segments:
            if isinstance(segment, str):
                segment = (segment, {})
            text, style = segment
            run = paragraph.add_run(text)
            set_run_font(run, style.get("size", 12), bold=style.get("bold", False), italic=style.get("italic", False))
            if style.get("sub"):
                run.font.subscript = True
        fmt(paragraph, align, after=after, keep=keep, indent=indent, first=first)
        return paragraph

    def bullet(self, *segments, level=0):
        paragraph = self.para(*segments, after=3, indent=0.9 + 0.7 * level, first=-0.5)
        first = paragraph.runs[0]
        bullet = copy.deepcopy(first._element)
        paragraph._p.insert(paragraph._p.index(first._element), bullet)
        from docx.text.run import Run

        marker = Run(bullet, paragraph)
        marker.text = ("•" if level == 0 else "–") + "\t"
        set_run_font(marker, 12, bold=False, italic=False)
        tabs = paragraph.paragraph_format.tab_stops
        tabs.add_tab_stop(Cm(0.9 + 0.7 * level))
        return paragraph

    def numbered(self, number, lead, text):
        paragraph = self.para((f"{number}\t", {"bold": True}), (lead, {"bold": True}), text, after=4, indent=0.9, first=-0.9)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
        return paragraph

    def equation(self, *segments):
        paragraph = self.para(*segments, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        return paragraph

    def figure_block(self, image, width_in, number, caption):
        paragraph = self.doc.add_paragraph()
        fmt(paragraph, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2, spacing=1.0, keep=True)
        paragraph.add_run().add_picture(str(image), width=Inches(width_in))
        cap = self.doc.add_paragraph()
        run = cap.add_run(f"Figure {number}: ")
        set_run_font(run, 12, bold=True)
        run = cap.add_run(caption)
        set_run_font(run, 12, italic=True)
        fmt(cap, WD_ALIGN_PARAGRAPH.CENTER, before=2, after=10, spacing=1.5)

    def _new_section(self, landscape: bool):
        from docx.enum.section import WD_ORIENT

        previous = self.doc.sections[-1]
        width, height = previous.page_width, previous.page_height
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        short, long = min(width, height), max(width, height)
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = (long, short) if landscape else (short, long)
        for name in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(section, name, Cm(2.54))
        numbering = section._sectPr.find(qn("w:pgNumType"))
        if numbering is not None and numbering.get(qn("w:start")) is not None:
            del numbering.attrib[qn("w:start")]
        return section

    def landscape_figures(self, figures):
        from PIL import Image

        self._new_section(landscape=True)
        for index, (image, number, caption) in enumerate(figures):
            with Image.open(image) as picture:
                aspect = picture.width / picture.height
            width = min(9.0, 5.7 * aspect)
            paragraph = self.doc.add_paragraph()
            fmt(paragraph, WD_ALIGN_PARAGRAPH.CENTER, before=0, after=2, spacing=1.0, keep=True)
            if index:
                paragraph.paragraph_format.page_break_before = True
            paragraph.add_run().add_picture(str(image), width=Inches(width))
            cap = self.doc.add_paragraph()
            run = cap.add_run(f"Figure {number}: ")
            set_run_font(run, 12, bold=True)
            run = cap.add_run(caption)
            set_run_font(run, 12, italic=True)
            fmt(cap, WD_ALIGN_PARAGRAPH.CENTER, before=4, after=0, spacing=1.5)
        self._new_section(landscape=False)

    def table(self, number, caption, header, rows, widths_cm, font_size=11):
        cap = self.doc.add_paragraph()
        run = cap.add_run(f"Table {number}: ")
        set_run_font(run, 12, bold=True)
        run = cap.add_run(caption)
        set_run_font(run, 12, italic=True)
        fmt(cap, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=4, spacing=1.5, keep=True)
        table = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = OxmlElement(f"w:{side}")
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), "6")
            node.set(qn("w:space"), "0")
            node.set(qn("w:color"), "404040")
            borders.append(node)
        table._tbl.tblPr.append(borders)
        for row_index, values in enumerate([header] + rows):
            row = table.rows[row_index]
            if row_index == 0:
                trpr = row._tr.get_or_add_trPr()
                repeat = OxmlElement("w:tblHeader")
                repeat.set(qn("w:val"), "true")
                trpr.append(repeat)
            for col, value in enumerate(values):
                cell = row.cells[col]
                cell.width = Cm(widths_cm[col])
                cell.paragraphs[0].text = ""
                paragraph = cell.paragraphs[0]
                run = paragraph.add_run(value)
                set_run_font(run, font_size, bold=row_index == 0)
                fmt(paragraph, WD_ALIGN_PARAGRAPH.LEFT, after=0, spacing=1.15)
                if row_index == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:val"), "clear")
                    shading.set(qn("w:color"), "auto")
                    shading.set(qn("w:fill"), "D9E2F3")
                    cell._tc.get_or_add_tcPr().append(shading)
        spacer = self.doc.add_paragraph()
        fmt(spacer, after=4, spacing=1.0)
        return table


def replace_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        run = paragraph.add_run(text)
        set_run_font(run)
        return
    runs[0].text = text
    for run in runs[1:]:
        run._element.getparent().remove(run._element)


def prepare_cover(doc: Document) -> None:
    paragraphs = doc.paragraphs
    texts = [p.text for p in paragraphs]
    end = next(i for i, t in enumerate(texts) if t.strip().startswith("Bhauti"))
    replace_text(paragraphs[2], "CreatorLens")
    replace_text(paragraphs[12], "(Project ID: [PLACEHOLDER — fill in later])")
    replace_text(paragraphs[14], "[Student Name — PLACEHOLDER] ([Roll Number — PLACEHOLDER])")
    replace_text(paragraphs[17], "[Supervisor Name — PLACEHOLDER]")
    replace_text(paragraphs[18], "([Supervisor Designation — PLACEHOLDER])")
    replace_text(paragraphs[21], "Pranveer Singh Institute of Technology")
    replace_text(paragraphs[22], "Kanpur – Agra – Delhi National Highway - 19")
    replace_text(paragraphs[23], "Bhauti – Kanpur – 209305")

    body = doc.element.body
    keep = {paragraph._p for paragraph in paragraphs[: end + 1]}
    for child in list(body.iterchildren()):
        if child.tag == qn("w:sectPr"):
            continue
        if child not in keep:
            body.remove(child)

    for paragraph in paragraphs[: end + 1]:
        for run in paragraph.runs:
            set_run_font(run)

    body_sectpr = doc.sections[0]._sectPr
    cover_sectpr = copy.deepcopy(body_sectpr)
    for child in list(cover_sectpr):
        if child.tag in (qn("w:headerReference"), qn("w:footerReference"), qn("w:pgNumType"), qn("w:titlePg")):
            cover_sectpr.remove(child)
    last = paragraphs[end]._p
    ppr = last.get_or_add_pPr()
    ppr.append(cover_sectpr)

    page_numbers = body_sectpr.find(qn("w:pgNumType"))
    if page_numbers is None:
        page_numbers = OxmlElement("w:pgNumType")
        body_sectpr.append(page_numbers)
    page_numbers.set(qn("w:start"), "1")
    for child in list(body_sectpr):
        if child.tag == qn("w:type"):
            body_sectpr.remove(child)

    section = doc.sections[-1]
    for name in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, name, Cm(2.54))
    section.different_first_page_header_footer = False
    doc.settings.odd_and_even_pages_header_footer = False
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    begin = paragraph.add_run()
    set_run_font(begin, 12)
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), "begin")
    begin._element.append(node)
    instr = paragraph.add_run()
    set_run_font(instr, 12)
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = " PAGE "
    instr._element.append(text)
    sep = paragraph.add_run()
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), "separate")
    sep._element.append(node)
    value = paragraph.add_run("1")
    set_run_font(value, 12)
    end_run = paragraph.add_run()
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), "end")
    end_run._element.append(node)


def normalise_styles(doc: Document) -> None:
    for name in ("Normal", "Title", "Heading 1", "List Paragraph"):
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        style.font.name = FONT
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attribute), FONT)
        for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if fonts.get(qn(attribute)) is not None:
                del fonts.attrib[qn(attribute)]
        if name == "Normal":
            style.font.size = Pt(12)


def write_content(b: Builder) -> None:
    # 1. Introduction ------------------------------------------------------------------
    b.heading("1. Introduction", page_break=False)
    b.para(
        "Content creators on platforms such as YouTube, Instagram and TikTok receive per-post counts of views, likes, "
        "comments and shares. Each number is easy to read on its own, but it is difficult to compare many posts that "
        "differ in age, platform and format. Older posts keep collecting views, a single viral post distorts an "
        "average, and rates calculated in different ways cannot be compared. As a result, creators often judge "
        "performance from a recent spike or from intuition instead of from a consistent analysis of their own data."
    )
    b.para(
        "CreatorLens is a web application developed as a B.Tech Computer Science and Engineering mini project in the "
        "creator-analytics domain. It converts a creator's own data into consistent metrics, charts and "
        "evidence-labelled insights. One view-based engagement rate, (likes + comments + shares) / views × 100, is "
        "used throughout, so that different posts and categories are measured in the same way. Missing counts are never "
        "filled in, sample data is always labelled as demo data, and every insight states the evidence behind it."
    )
    b.para(
        "The implemented system uses a React and TypeScript frontend built with Vite ",
        cite("react", "typescript", "vite"),
        " and a FastAPI backend ",
        cite("fastapi"),
        ". Metrics are calculated with Pandas and NumPy ",
        cite("pandas", "numpy"),
        ", and results are stored in local SQLite databases ",
        cite("sqlite"),
        ". Public YouTube video metadata is read with yt-dlp ",
        cite("ytdlp"),
        ". Two prediction features are included: an engagement-rate model that compares a mean baseline, a decision "
        "tree, a random forest and XGBoost using scikit-learn ",
        cite("sklearn", "xgboost"),
        ", and a view forecast that uses Holt's linear exponential smoothing ",
        cite("holt", "statsmodels"),
        " only when it performs better than a last-value baseline. Google Gemini ",
        cite("gemini"),
        " is optional and is used only to word summaries; all numbers come from the application's own calculations.",
    )
    b.para(
        "The scope is limited to data that the creator supplies (CSV, PDF and DOCX uploads, and dated view history) and "
        "to public YouTube metadata. Private platform analytics such as impressions, click-through rate and audience "
        "retention are not available because creator OAuth access is not connected. CreatorLens does not publish or "
        "schedule posts, and it does not guarantee growth. Its relevance lies in giving small and early-stage creators a "
        "transparent, reproducible way to understand past performance and to estimate future results with stated "
        "limitations."
    )

    # 2. Project Objective ----------------------------------------------------------------
    b.heading("2. Project Objective")
    b.para(
        "The objectives below are stated so that each can be checked against the working application and its "
        "automated tests. Table 2.1 separates what is implemented from what is partial or planned."
    )
    objectives = [
        ("Dataset upload and validation: ", "accept a creator's post data as a CSV file (and PDF or DOCX documents up to "
         "15 MB), reject unsupported or malformed files with a specific message, and save each successful analysis "
         "together with the original file."),
        ("Consistent metrics: ", "calculate the view-based engagement rate for every valid post, together with category, "
         "platform and weekly-trend summaries and outlier flags, without treating a missing count as zero."),
        ("Visual insights: ", "present the results as interactive charts and short insights on the Dashboard, Analytics, "
         "Insights and saved-analysis pages, each with the evidence and limitations behind it."),
        ("YouTube video analysis: ", "analyse a public YouTube video from its URL (views, likes, comments, duration, "
         "engagement rate, views per day) and generate an evidence-based performance diagnosis with an action plan."),
        ("Engagement-rate prediction: ", "predict the engagement rate of a planned post from fields known before "
         "publication, select the best of four models on held-out posts, and explain which inputs influenced the result."),
        ("View forecasting: ", "forecast a video's views after 24 hours, 3 days and 7 days from dated observations, using "
         "Holt's model only when back-testing shows that it is at least 10% more accurate than a last-value baseline."),
        ("Creator-specific records: ", "keep saved analyses associated with an account through server-side ownership "
         "checks (implemented in the backend; sign-in is currently paused, as noted in Table 2.1)."),
    ]
    for index, (lead, text) in enumerate(objectives, start=1):
        b.numbered(f"O{index}.", lead, text)
    b.table(
        "2.1",
        "Implementation status of CreatorLens features (verified from the repository on 21 Oct 2026)",
        ["Feature", "Status", "Evidence in repository"],
        [
            ["Demo dashboard, analytics and insights (16 labelled demo posts)", "Implemented", "demo_dataset.py; /api/dashboard/*"],
            ["CSV/PDF/DOCX upload, validation, saved analysis and questions", "Implemented", "document_analysis.py; /api/documents"],
            ["YouTube public-video analysis and AI Performance Diagnosis", "Implemented", "youtube_analysis.py; diagnosis.py"],
            ["Roast My Video and Content Court (creative feedback)", "Implemented", "roast.py"],
            ["Creator analytics CSV (dated rows) analysis", "Implemented", "analytics_analysis.py"],
            ["Engagement-rate prediction model", "Implemented; evaluated only on a synthetic test dataset so far", "engagement_model.py; /api/predictions/*"],
            ["View forecasting (Holt vs. last-value baseline)", "Implemented; saved videos currently have one observation each, so no real forecast yet", "prediction.py; view_history.py"],
            ["Accounts: registration, email verification, login, password reset", "Partially implemented: backend complete, sign-in paused; email needs SMTP", "auth.py; auth_routes.py"],
            ["Gemini-written summaries and wording", "Optional: needs API key and quota; calculated results remain without it", "diagnosis.py; document_analysis.py"],
            ["Gemini video-file analysis", "Partial: backend endpoint only, no page", "video_routes.py (/api/videos/analyze)"],
            ["Private YouTube Analytics (impressions, CTR, retention)", "Planned: OAuth not connected", "youtube_client.py"],
            ["Automatic scheduled view snapshots", "Planned: manual or command-line collection only", "view_history.py --collect"],
            ["PostgreSQL and Clerk authentication", "Not implemented (planning documents only)", "TRD.md; architecture.md"],
        ],
        [6.0, 5.2, 4.9],
        font_size=10.5,
    )
    b.para(
        ("Expected benefit to the end user: ", {"bold": True}),
        "a creator can upload existing data and immediately see which content types and platforms earn more engagement "
        "per view, understand why a video looks strong or weak from the evidence available, and obtain a prediction or "
        "forecast that is accompanied by its measured error and limitations. This reduces manual spreadsheet work and "
        "discourages decisions based on a single post.",
    )

    # 3. Feasibility Study ------------------------------------------------------------------
    b.heading("3. Feasibility Study")
    b.para(
        "A feasibility study examines whether the project can succeed by assessing its technical, operational, economic, "
        "schedule and legal aspects. Each aspect is assessed below against the system that has actually been built."
    )
    b.subheading("3.1 Technical Feasibility")
    b.para(
        "All core technologies are mature and open source: React, TypeScript and Vite for the interface ",
        cite("react", "typescript", "vite"),
        ", Recharts for charts ",
        cite("recharts"),
        ", Tailwind CSS for styling ",
        cite("tailwind"),
        ", FastAPI for the REST API ",
        cite("fastapi"),
        ", and Pandas, NumPy, scikit-learn, XGBoost and statsmodels for analysis and modelling ",
        cite("pandas", "numpy", "sklearn", "xgboost", "statsmodels"),
        ". The datasets are small and tabular, so training and forecasting run on an ordinary laptop CPU without a "
        "GPU. SQLite ",
        cite("sqlite"),
        " removes the need for a separate database server. The backend has 142 automated pytest tests ",
        cite("pytest"),
        ", all of which passed on 21 October 2026, and the frontend passes a TypeScript type check and a production build.",
    )
    b.para(
        "Two external dependencies carry technical risk. First, yt-dlp ",
        cite("ytdlp"),
        " depends on YouTube's public page structure and can be blocked by bot checks; the backend reports such failures "
        "as clear errors instead of inventing counts. Second, the optional Gemini API ",
        cite("gemini"),
        " is limited by quota; during development the free-tier daily request limit was reached, and the application "
        "continued with its calculated results. The project is therefore technically feasible with these documented "
        "limitations.",
    )
    b.subheading("3.2 Operational Feasibility")
    b.para(
        "A creator uses CreatorLens in a browser without writing code. The implemented workflow is: open the Dashboard "
        "to explore labelled demo data; upload a posts CSV on the Upload Dataset page and read the calculated summary, "
        "charts and outliers; paste a public YouTube URL on the Analyze YouTube page to obtain statistics, a diagnosis "
        "and optional creative feedback; and open the Predictions page to train the engagement-rate model on a saved "
        "posts CSV or to request a view forecast for a saved video. Every page explains when data is insufficient, for "
        "example that one public snapshot cannot establish a trend, and tells the creator what to collect next. The "
        "main operational limitation is that view forecasts need dated observations collected over time or imported "
        "from a history CSV."
    )
    b.subheading("3.3 Economic Feasibility")
    b.para(
        "Development uses free and open-source software on an existing computer, so no licence or server cost has been "
        "incurred. Table 3.1 summarises the cost position. Gemini is the only component that may require payment "
        "beyond its free usage limits; it is optional, and its cost was not estimated for this synopsis. Hosting costs "
        "will arise only if the application is deployed publicly in future."
    )
    b.table(
        "3.1",
        "Cost assessment",
        ["Item", "Cost for this project", "Remarks"],
        [
            ["Frontend, backend and ML libraries", "Free (open source)", "React, FastAPI, Pandas, scikit-learn, XGBoost, statsmodels"],
            ["Database", "Free", "SQLite files on the local disk; no database server"],
            ["Development tools", "Free", "Python, Node.js, npm, Git, code editor"],
            ["Google Gemini API", "Free tier with limits; paid use is usage-based", "Optional; the app works without it"],
            ["Public YouTube metadata", "Free", "Read with yt-dlp; no paid API key"],
            ["Deployment and hosting", "Not incurred", "Future cost depends on the chosen provider"],
        ],
        [4.4, 4.6, 7.1],
    )
    b.subheading("3.4 Schedule Feasibility")
    b.para(
        "Git commits and file timestamps show that the planning documents were written on 14 October 2026 and that the "
        "implemented modules were developed between 14 and 20 October 2026. The remaining work, namely integration "
        "testing, re-enabling sign-in, evaluation on real creator data, and the final report, is shown as a proposed "
        "plan with estimated dates in the Gantt chart (Figure 4.4 and Table 4.1 in Section 4.4). The project is "
        "feasible within the mini-project schedule because the core features already work and the remaining tasks are "
        "testing, evaluation and documentation."
    )
    b.subheading("3.5 Legal and Ethical Feasibility")
    b.para(
        "CreatorLens processes data that creators upload about their own content. Uploaded files and databases are kept "
        "in the local backend/data folder, which is excluded from version control. Account passwords are stored as "
        "PBKDF2-SHA256 hashes with 600,000 iterations ",
        cite("rfc8018", "owasp"),
        ", sessions use HttpOnly cookies, and state-changing requests require a CSRF token. Because the application "
        "handles personal data such as names and email addresses, any public deployment in India would need to "
        "consider the Digital Personal Data Protection Act, 2023 ",
        cite("dpdp"),
        "; compliance has not been assessed and is not claimed.",
    )
    b.para(
        "Public YouTube metadata is read with yt-dlp for academic demonstration. YouTube's Terms of Service ",
        cite("youtube_tos"),
        " restrict automated access to the service, so a production version should use YouTube's official APIs with "
        "creator consent. When a Gemini key is configured, extracted document text and calculated evidence are sent to "
        "Google ",
        cite("gemini"),
        ", and users should be informed of this. Ethically, the system never fabricates metrics, labels demo data, "
        "presents predictions as estimates rather than guarantees, describes feature importance as association rather "
        "than cause, and keeps creative feedback focused on content instead of the person.",
    )

    # 4. Methodology --------------------------------------------------------------------------
    b.heading("4. Methodology / Planning of Work")
    b.para(
        "The following steps describe how CreatorLens processes data from input to result. Each step reflects the code "
        "in the repository; partial or planned elements are identified explicitly."
    )
    steps = [
        ("User access. ", "The creator opens the React application in a browser. Account registration, email "
         "verification, login, logout and password reset are implemented in the backend, but sign-in is currently "
         "paused, so all saved work is stored under a single local account."),
        ("Data input. ", "The creator can upload a CSV, PDF or DOCX file (up to 15 MB) on the Upload Dataset page, "
         "paste a public YouTube URL, import a dated view-history CSV with the columns observed_at and views, or "
         "explore the built-in demo dataset of 16 labelled posts."),
        ("Validation and preprocessing. ", "The backend checks file extension, content type and size; parses CSV as "
         "UTF-8 within row and column limits; converts numeric fields; drops invalid rows with a stated reason; "
         "converts timestamps to UTC and rejects duplicates and future dates. PDF and DOCX text is extracted with pypdf "
         "and python-docx. A YouTube URL is reduced to its video ID, and public metadata is cached for six hours."),
        ("Analysis and metric calculation. ", "Pandas computes the engagement rate for each valid post, category and "
         "platform averages, weekly trends, and outliers outside 1.5 times the interquartile range. For a YouTube "
         "video, views per day and the comment rate are calculated, and a creator median is used as a baseline only "
         "when at least three other videos from the same channel have been saved."),
        ("Visualisation and insight generation. ", "Recharts draws trend, comparison and interaction charts. The "
         "performance diagnosis is written from rule-based evidence and labelled Limited, Moderate or Stronger "
         "evidence. Gemini may reword it only if every cited figure still matches the calculation."),
        ("Prediction and forecasting. ", "The engagement-rate model and the view forecast are run on request from the "
         "Predictions page. Both are described after these steps."),
        ("Storage and retrieval. ", "Analyses are saved as JSON payloads in SQLite tables with an owner column; "
         "uploaded files are kept in backend/data/files and trained models in backend/data/models (Figure 4.3)."),
        ("Display of results. ", "The frontend shows the results on the Dashboard, Analytics, Insights, Saved "
         "analyses, YouTube analysis and Predictions pages, including loading states and clear error messages."),
    ]
    for index, (lead, text) in enumerate(steps, start=1):
        b.numbered(f"Step {index}.", lead, text)

    b.subheading("Engagement-Rate Prediction Model (Step 6)")
    b.para(
        "The target is the engagement rate of a post. To avoid target leakage, the model uses only fields that are "
        "known before publication: platform, category, posting date (weekday and month), posting hour, title (length, "
        "word count, digits, question mark) and duration. Views, likes, comments, shares and other outcome columns are "
        "excluded from the inputs. A saved CSV must contain at least 30 usable posts with views above zero."
    )
    b.para(
        "The posts are sorted by date, and the latest 20% (at least six posts) are held out for testing. Four "
        "scikit-learn pipelines are trained on the earlier posts: a mean baseline, a decision tree ",
        cite("cart"),
        ", a random forest ",
        cite("rf"),
        " and XGBoost ",
        cite("xgboost"),
        ". The model with the lowest mean absolute error (MAE) on the held-out posts is selected and refitted on all "
        "posts; if no model beats the mean baseline, predictions are refused. Input importance is measured by "
        "shuffling each input on the held-out posts, and when XGBoost is selected, per-prediction TreeSHAP "
        "contributions ",
        cite("treeshap"),
        " are shown. On a 200-post synthetic test file, XGBoost achieved an MAE of 0.84 percentage points against 0.92 "
        "for the mean baseline, 0.89 for the random forest and 1.03 for the decision tree. These figures come from "
        "test data and are not evidence of accuracy on real creator data.",
    )
    b.subheading("View Forecasting with Holt's Linear Exponential Smoothing (Step 6)")
    b.para(
        "A video's cumulative view count is forecast from dated observations of the same video. Holt's method ",
        cite("holt", "fpp"),
        " maintains a smoothed level ℓ and trend b, which are updated at each observation y:",
    )
    b.equation("ℓ", ("t", {"sub": True}), " = α·y", ("t", {"sub": True}), " + (1 − α)(ℓ", ("t−1", {"sub": True}),
               " + b", ("t−1", {"sub": True}), ")")
    b.equation("b", ("t", {"sub": True}), " = β(ℓ", ("t", {"sub": True}), " − ℓ", ("t−1", {"sub": True}), ") + (1 − β)·b",
               ("t−1", {"sub": True}), ",     ŷ", ("t+h", {"sub": True}), " = ℓ", ("t", {"sub": True}), " + h·b", ("t", {"sub": True}))
    b.para(
        "Here α and β are smoothing parameters estimated by statsmodels ",
        cite("statsmodels"),
        ", and h is the number of steps ahead. Because real snapshots are taken at irregular times, the views are first "
        "placed on a regular 1-, 3-, 6-, 12- or 24-hour grid by linear interpolation between real observations. The "
        "baseline predicts that future views equal the latest count, ŷ = y",
        ("t", {"sub": True}),
        ".",
    )
    b.para(
        "With fewer than three observations, no forecast is produced. With three to five, a provisional baseline "
        "forecast is shown. With six or more, a rolling-origin back-test ",
        cite("tashman"),
        " is run: from each earlier origin, both methods predict the real observations that follow within seven days, "
        "and Holt's model is selected only if its MAE is at least 10% lower than the baseline's. Forecasts for 24 "
        "hours, 3 days and 7 days are never below the latest count, and any horizon longer than twice the recorded "
        "history is hidden. Saved videos currently have one observation each, so the application shows guidance to "
        "collect snapshots instead of a forecast.",
    )

    b.subheading("4.1 System Architecture Diagram")
    b.para(
        "Figure 4.1 shows the verified architecture. The browser runs the React single-page application, which calls "
        "the FastAPI backend through the /api path. API routers delegate work to service modules for data processing, "
        "document extraction, diagnosis, the engagement model and view forecasting. Services read and write SQLite "
        "databases and local folders, read public YouTube metadata with yt-dlp, and optionally call the Gemini API. "
        "Figures 4.1 to 4.4 are placed on landscape pages at the end of this section so that their labels remain legible."
    )
    b.subheading("4.2 Data Flow Diagram")
    b.para(
        "Figure 4.2 is a level-1 data flow diagram. The creator supplies files, URLs, planned-post details and history "
        "CSVs to five processes. The processes exchange data with four data stores and with the two external entities, "
        "YouTube public pages and the optional Gemini API, and they return metrics, diagnoses, predictions and forecasts "
        "to the creator."
    )
    b.subheading("4.3 ER Diagram")
    b.para(
        "An ER diagram was chosen because CreatorLens persists its data in a relational SQLite schema. Figure 4.3 shows "
        "the ten tables defined in the code, grouped by database file. Enforced foreign keys link sessions and "
        "email tokens to users and questions to documents. The user_id columns of the analysis tables and the video_id "
        "links between YouTube tables are logical links without foreign-key constraints, as implemented."
    )
    b.subheading("4.4 Project Timeline and Gantt Chart")
    b.para(
        "Figure 4.4 shows the project timeline in days. Completed phases are dated from Git commits (15 and 18 "
        "October 2026) and from file creation and modification times (14 to 20 October 2026). Work after the synopsis "
        "date of 21 October 2026 is a proposed plan; its dates are estimates and will be confirmed with the supervisor. "
        "Table 4.1 lists the same phases with their status and evidence."
    )
    b.table(
        "4.1",
        "Project phases, dates and status",
        ["Phase", "Dates (2026)", "Status", "Evidence"],
        [
            ["Requirement analysis and problem definition", "14 Oct", "Completed", "PRD.md, TRD.md"],
            ["System design and architecture", "14 Oct", "Completed", "architecture.md"],
            ["Frontend shell, demo dashboard, FastAPI foundation", "14–15 Oct", "Completed", "Commit 5a4dc11"],
            ["CSV validation and document analysis", "15 Oct", "Completed", "Commit d3d28ec"],
            ["YouTube public-metadata analysis", "15–16 Oct", "Completed", "Commit d3d28ec; youtube_client.py"],
            ["Creator analytics CSV and diagnosis", "16 Oct", "Completed", "analytics_csv.py; diagnosis.py"],
            ["Accounts and SQLite data ownership", "17 Oct", "Completed; sign-in now paused", "auth.py; commit c9aca1b"],
            ["Roast My Video", "17 Oct", "Completed", "roast.py; commit c9aca1b"],
            ["Engagement-rate model", "18 Oct", "Completed (not yet committed)", "engagement_model.py"],
            ["View forecasting", "19–20 Oct", "Completed (not yet committed)", "prediction.py; view_history.py"],
            ["Automated backend tests", "14–20 Oct", "Completed (142 tests)", "backend/tests"],
            ["Literature review and synopsis", "21 Oct", "In progress", "This document"],
            ["Integration testing and debugging", "22–31 Oct (est.)", "Proposed", "—"],
            ["Re-enable sign-in, verify per-user isolation", "27 Oct–5 Nov (est.)", "Proposed", "—"],
            ["Collect real creator data and evaluate models", "27 Oct–9 Nov (est.)", "Proposed", "—"],
            ["Final report, presentation and submission", "6–14 Nov (est.)", "Proposed", "—"],
        ],
        [6.2, 3.3, 3.6, 3.0],
        font_size=10.5,
    )
    b.landscape_figures([
        (OUT_DIR / "CreatorLens_Architecture_Diagram.png", "4.1", "System architecture of CreatorLens (implemented components)"),
        (OUT_DIR / "CreatorLens_Data_Flow_Diagram.png", "4.2", "Level-1 data flow diagram of CreatorLens"),
        (OUT_DIR / "CreatorLens_ER_Diagram.png", "4.3", "Entity-relationship diagram of the implemented SQLite schema"),
        (OUT_DIR / "CreatorLens_Gantt_Chart.png", "4.4", "Gantt chart of the CreatorLens project timeline"),
    ])

    # 5. Tools / Technology -------------------------------------------------------------------
    b.heading("5. Tools/Technology Used")
    b.subheading("5.1 Minimum Hardware Requirements")
    b.para(
        "The hardware below is required for developing and running the project locally. The specifications are "
        "estimated minimums; only the disk usage figures were measured in the repository. The exact hardware of the "
        "development computer was not recorded in the project files."
    )
    b.table(
        "5.1",
        "Minimum hardware requirements (estimated)",
        ["Component", "Minimum requirement", "Remarks"],
        [
            ["CPU", "64-bit dual-core x86-64 processor", "Quad-core recommended for running the frontend, backend and browser together"],
            ["RAM", "4 GB (8 GB recommended)", "Model training uses small tabular files in memory"],
            ["GPU", "Not required", "XGBoost, scikit-learn and statsmodels run on the CPU in this project"],
            ["HDD/SSD", "At least 3 GB free; SSD recommended", "Measured: frontend node_modules ≈ 174 MB, backend virtual environment ≈ 527 MB"],
            ["Others", "Internet connection; modern web browser", "Needed for installing packages, reading YouTube metadata and optional Gemini calls"],
        ],
        [2.6, 5.0, 8.5],
    )
    b.subheading("5.2 Minimum Software Requirements")
    b.para(
        "Table 5.2 lists the software verified in the repository's dependency files, installed packages and source code. "
        "Versions are those installed in the development environment on 21 October 2026."
    )
    b.table(
        "5.2",
        "Software and technologies used (verified)",
        ["Category", "Technology and version", "Role in CreatorLens"],
        [
            ["OS", "Windows 11 Home (development)", "Development and testing environment"],
            ["Runtime", "Python 3.14.4; Node.js 24.15.0 with npm 11.12.1", "Backend runtime; frontend tooling"],
            ["Frontend", "React 19.3.1, React DOM 19.3.1, TypeScript 6.0.3", "User interface and type safety"],
            ["Build tool", "Vite 8.3.4 with @vitejs/plugin-react 6.1.2", "Development server, /api proxy and production build"],
            ["Routing and UI", "React Router DOM 7.18.4; Tailwind CSS 4.3.3; Lucide React 1.57.0", "Page routing, styling and icons"],
            ["Charts", "Recharts 3.10.1", "Trend, comparison, prediction and forecast charts"],
            ["Backend", "FastAPI 0.143.1; Uvicorn 0.54.0; Pydantic 2.14.1; python-multipart 0.0.32", "REST API, ASGI server, request and response validation, file uploads"],
            ["Data analysis", "Pandas 2.3.3; NumPy 2.5.4", "CSV parsing, metrics, trends and outliers"],
            ["Machine learning", "scikit-learn 1.9.1; XGBoost 3.4.1; joblib 1.6.0", "Engagement-rate model pipelines, model comparison and model storage"],
            ["Time-series forecasting", "statsmodels 0.15.0 (Holt's linear exponential smoothing)", "View forecasting with back-testing against a last-value baseline"],
            ["Database", "SQLite 3.50.4 (Python sqlite3 module)", "youtube, documents, analytics and auth databases"],
            ["Document extraction", "pypdf 6.20.1; python-docx 1.2.0", "Text extraction from PDF and DOCX uploads"],
            ["External data", "yt-dlp 2026.8.19", "Reading public YouTube video metadata"],
            ["AI service (optional)", "google-genai 2.30.1; default model gemini-3.8-flash", "Server-side summaries and wording when a key is configured"],
            ["Authentication", "Custom: PBKDF2-SHA256, HttpOnly session cookie, CSRF token", "Implemented in the backend; sign-in currently paused"],
            ["Testing and quality", "pytest 8.4.2; httpx 0.28.1; oxlint 1.88.0", "142 backend tests; API test client; frontend lint"],
            ["Development tools", "Git; Cursor code editor; web browser", "Version control, editing and manual testing"],
        ],
        [3.2, 6.3, 6.6],
        font_size=10.5,
    )
    b.para(
        ("Planned but not used: ", {"bold": True}),
        "the planning documents propose PostgreSQL with SQLAlchemy and Alembic, and Clerk authentication. None of these "
        "is used in the current code; the implementation uses SQLite and its own session-based authentication.",
    )

    # 6. References -----------------------------------------------------------------------------
    b.heading("6. References")
    for index, key in enumerate(ORDER, start=1):
        paragraph = b.para(f"[{index}]\t{REFERENCES[key]}", align=WD_ALIGN_PARAGRAPH.LEFT, after=4, indent=1.0, first=-1.0)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))


def main() -> None:
    doc = Document(str(TEMPLATE))
    normalise_styles(doc)
    prepare_cover(doc)
    builder = Builder(doc)
    write_content(builder)
    unused = [key for key in REFERENCES if key not in ORDER]
    if unused:
        raise SystemExit(f"References defined but not cited: {unused}")
    core = doc.core_properties
    core.title = "CreatorLens – Mini Project Synopsis"
    core.subject = "B.Tech CSE Mini Project Synopsis"
    core.author = "[Student Name — PLACEHOLDER]"
    doc.save(str(OUTPUT))
    print("saved", OUTPUT, "references:", len(ORDER))


if __name__ == "__main__":
    main()
