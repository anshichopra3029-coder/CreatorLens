"""Extract text from PDF and DOCX files.

The extracted text is data. It is never executed.
"""

from __future__ import annotations

import io

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader


class ExtractError(Exception):
    def __init__(self, status_code: int, detail: str) -> None:
        self.status_code = status_code
        self.detail = detail
        super().__init__(detail)


def extract_pdf(raw: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(raw))
    except Exception as exc:
        raise ExtractError(400, "The PDF could not be read. It may be corrupted.") from exc
    if reader.is_encrypted:
        raise ExtractError(400, "This PDF is encrypted. Upload a PDF without a password.")
    parts: list[str] = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception as exc:
            raise ExtractError(400, "The PDF could not be read. It may be corrupted.") from exc
    return "\n".join(parts).strip()


def extract_docx(raw: bytes) -> str:
    try:
        document = Document(io.BytesIO(raw))
    except (PackageNotFoundError, ValueError) as exc:
        raise ExtractError(400, "The DOCX file could not be read. It may be corrupted.") from exc
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()
