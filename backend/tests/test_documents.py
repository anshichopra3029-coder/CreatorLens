"""Document analysis tests. Narrative calls use a stand-in client, not Gemini."""

import json
from types import SimpleNamespace

import pytest
from docx import Document

from app.core.config import settings
from app.services.document_analysis import analyze_upload
from app.services.document_analysis import MISSING_DOCUMENT_KEY

CSV = (
    b"post_id,platform,category,post_datetime,views,likes,comments,shares,followers_gained\n"
    b"p1,Instagram,Tutorial,2026-10-01,1000,50,10,5,2\n"
    b"p2,YouTube,Tutorial,2026-10-08,2000,100,20,10,3\n"
    b"p3,TikTok,Vlog,2026-10-08,500,10,0,0,1\n"
)

NARRATIVE = {
    "subject": "Tutorial and vlog posts",
    "summary": "Three posts are in the file. The calculated average engagement rate is 5.0.",
    "main_points": ["Tutorial posts share an average rate of 6.5."],
    "key_findings": ["Total interactions across rated posts are 205."],
    "facts": ["The average engagement rate is 5.0."],
    "dates": ["2026-10-01"],
    "entities": ["Tutorial"],
    "risks": [],
    "recommendations": ["Compare the Vlog post with the Tutorial posts."],
}


class _Interactions:
    def __init__(self, output: dict) -> None:
        self.output = output
        self.calls = 0

    def create(self, **kwargs: object) -> SimpleNamespace:
        self.calls += 1
        assert kwargs["model"] == "gemini-3.8-flash"
        return SimpleNamespace(output_text=json.dumps(self.output))


class _Client:
    def __init__(self, output: dict) -> None:
        self.interactions = _Interactions(output)


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    monkeypatch.setattr(settings, "gemini_api_key", "")
    monkeypatch.setattr(settings, "gemini_model", "gemini-3.8-flash")
    return tmp_path


def _post(client, name: str, content: bytes, content_type: str):
    return client.post("/api/documents", files={"file": (name, content, content_type)})


def test_csv_calculates_engagement_and_persists_without_gemini(member, isolated):
    response = _post(member, "posts.csv", CSV, "text/csv")
    assert response.status_code == 200
    body = response.json()
    assert body["file_saved"] is True
    assert body["is_demo"] is False
    assert body["analysis_performed"] is True
    assert body["narrative_status"] == "unavailable"
    assert body["summary"] is None
    profile = body["csv_profile"]
    assert profile["row_count"] == 3
    assert profile["engagement_available"] is True
    assert profile["avg_engagement_rate"] == 5.0
    assert profile["total_interactions"] == 205
    assert profile["best_category"] == "Tutorial"
    assert profile["posts_analyzed"] == 3
    saved = member.get(f"/api/documents/{body['id']}")
    assert saved.status_code == 200
    assert saved.json()["csv_profile"]["avg_engagement_rate"] == 5.0
    assert (isolated / "files" / f"{body['id']}.csv").is_file()
    listed = member.get("/api/documents")
    assert listed.json()["count"] == 1


def test_csv_metadata_route_still_does_not_save(client, isolated):
    response = client.post("/api/upload", files={"file": ("posts.csv", CSV, "text/csv")})
    assert response.status_code == 200
    assert response.json()["analysis_performed"] is False
    assert response.json()["file_saved"] is False
    assert list(isolated.glob("files/*")) == []


def test_header_only_csv_is_rejected(member, isolated):
    response = _post(member, "empty.csv", b"post_id,views\n", "text/csv")
    assert response.status_code == 400
    assert member.get("/api/documents").json()["count"] == 0


def test_unsupported_file_is_rejected(member, isolated):
    response = _post(member, "notes.txt", b"hello", "text/plain")
    assert response.status_code == 415
    assert "not supported" in response.json()["detail"]


def test_docx_text_is_extracted_and_saved(member, isolated, tmp_path):
    path = tmp_path / "menu.docx"
    document = Document()
    document.add_paragraph("The spring menu launches on 12 May 2026. The budget is 40 dollars.")
    document.save(path)
    response = _post(
        member,
        "menu.docx",
        path.read_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert response.status_code == 200
    body = response.json()
    assert body["format"] == "docx"
    assert body["file_saved"] is True
    assert "40 dollars" in body["text_preview"]
    assert body["summary"] is None
    again = member.get(f"/api/documents/{body['id']}")
    assert "40 dollars" in again.json()["text_preview"]


def test_text_pdf_extracts_visible_words(member, isolated):
    response = _post(member, "budget.pdf", _text_pdf("Budget note spend 40 dollars."), "application/pdf")
    assert response.status_code == 200
    body = response.json()
    assert body["format"] == "pdf"
    assert body["file_saved"] is True
    assert body["status"] == "extracted"
    assert "40 dollars" in body["text_preview"]
    assert body["summary"] is None
    again = member.get(f"/api/documents/{body['id']}")
    assert "40 dollars" in again.json()["text_preview"]


def test_blank_pdf_is_saved_as_unreadable(member, isolated):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    raw = writer.to_bytes() if hasattr(writer, "to_bytes") else _writer_bytes(writer)
    response = _post(member, "scan.pdf", raw, "application/pdf")
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "unreadable"
    assert body["analysis_performed"] is False
    assert body["summary"] is None
    assert "No summary" in body["narrative_message"]
    assert (isolated / "files" / f"{body['id']}.pdf").is_file()


def test_corrupted_pdf_is_not_saved(member, isolated):
    response = _post(member, "broken.pdf", b"not a pdf", "application/pdf")
    assert response.status_code == 400
    assert member.get("/api/documents").json()["count"] == 0
    assert list((isolated / "files").glob("*")) == [] if (isolated / "files").exists() else True


def test_question_without_key_does_not_invent_an_answer(member, isolated, tmp_path):
    path = tmp_path / "note.docx"
    document = Document()
    document.add_paragraph("The workshop is in Mumbai on 2 June 2026.")
    document.save(path)
    created = _post(
        member,
        "note.docx",
        path.read_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    document_id = created.json()["id"]
    answer = member.post(f"/api/documents/{document_id}/questions", json={"question": "Where is the workshop?"})
    assert answer.status_code == 503
    assert answer.json()["detail"] == MISSING_DOCUMENT_KEY
    stored = member.get(f"/api/documents/{document_id}")
    assert stored.json()["questions"] == []


def test_schema_valid_narrative_is_saved_without_calling_gemini(isolated):
    secret = "test-key-not-real"
    settings.gemini_api_key = secret
    fake = _Client(NARRATIVE)
    result = analyze_upload(CSV, filename="posts.csv", content_type="text/csv", user_id="test-user", client=fake)
    assert result.narrative_status == "generated"
    assert result.summary is not None
    assert "5.0" in result.summary
    assert result.csv_profile is not None
    assert result.csv_profile.avg_engagement_rate == 5.0
    assert fake.interactions.calls == 1
    assert secret not in result.model_dump_json()
    settings.gemini_api_key = ""


def _text_pdf(sentence: str) -> bytes:
    content = f"BT /F1 18 Tf 72 720 Td ({sentence}) Tj ET\n".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%sendstream" % (len(content), content),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    parts = [b"%PDF-1.4\n"]
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(sum(len(part) for part in parts))
        parts.append(f"{index} 0 obj\n".encode("ascii") + body + b"\nendobj\n")
    xref_pos = sum(len(part) for part in parts)
    xref = [b"xref\n", f"0 {len(objects) + 1}\n".encode("ascii"), b"0000000000 65535 f \n"]
    for offset in offsets[1:]:
        xref.append(f"{offset:010d} 00000 n \n".encode("ascii"))
    trailer = (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode("ascii")
    )
    return b"".join(parts + xref + [trailer])


def _writer_bytes(writer) -> bytes:
    import io

    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()
